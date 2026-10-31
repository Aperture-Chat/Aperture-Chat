"""Real LibreOffice against real hostile documents.

Every vector in ``crafted_documents`` is converted through the production path
and must neither put the synthetic marker on a page, nor touch the local canary,
nor run a macro. The control converts the same files on a stock profile with the
raw bytes, to show the traps are live: without it, a clean result could mean only
that a vector was broken. Skipped when LibreOffice is not installed.
"""

from __future__ import annotations

import http.server
import io
import os
import subprocess
import threading
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

import pikepdf
import pypdf
import pypdfium2
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.util import Inches

from app.core.document_rendition import (
    ConversionError,
    RenditionOptions,
    convert_to_pdf,
    find_soffice,
    prepare_input,
    profile_settings_xml,
    soffice_command,
    sniff_office_format,
    write_profile,
)
from app.core.knowledge_originals import DocumentRef, KnowledgeOriginalStore
from app.main import app
from app.models.schemas import (
    Group,
    KnowledgeConfig,
    KnowledgeDocument,
    KnowledgeDocumentWall,
    Role,
    User,
)
from app.repositories.deps import get_store
from tests.crafted_documents import (
    DRAWING,
    FIRST_PAGE_TEXT,
    NESTED_WORDS,
    OFFICE_RELS,
    SLIDE,
    VECTORS,
    Traps,
    Vector,
    inert_nested_document,
    linked_picture,
    masked_nested_document,
    ole_frame,
    retyped_chunk,
)
from tests.crafted_documents import docx as crafted_docx

SOFFICE = find_soffice()
IN_IMAGE = os.environ.get("APERTURE_RENDITION_IMAGE_PROOF") == "1"
pytestmark = pytest.mark.skipif(
    SOFFICE is None and not IN_IMAGE, reason="LibreOffice is not installed"
)

MARKER = "APERTUREMARKER7741"
MAGENTA = (255, 0, 255)
# Long enough for a busy machine: a timeout proves nothing and fails the test.
OPTIONS = RenditionOptions(
    soffice=SOFFICE, timeout_seconds=240.0, memory_mb=4096, require_isolation=False
)

# Stopped before or by the import, so there is no PDF to inspect. Everything else
# must convert, or a pipeline that always failed would look perfectly safe.
REFUSED = {
    "docx-entity-file",
    "docx-entity-http",
    "docx-altchunk",
    "docx-subdocument",
    "docx-masked-nested-document",
    "docx-html-chunk-file",
    "docx-html-chunk-http",
    "docx-retyped-chunk-file",
    "docx-retyped-chunk-http",
    "docx-unknown-type-chunk-file",
    "docx-utf16-chunk-file",
    "docx-alternate-content-chunk-file",
    "xlsx-entity-file",
    "odt-linked-image-file",
    "odt-linked-image-http",
    "odt-section-file",
    "odt-section-http",
    "odt-ole-link",
    "odt-floating-frame",
    "odt-font-face",
    "odt-background-image",
    "odt-entity-file",
    "odt-xforms-http",
    "odt-xforms-file",
    "odt-encoded-traversal",
    "odt-macro-on-load",
    "odt-nested-object-file",
    "ods-table-source",
    "ods-cell-range-link",
    "odp-linked-image-file",
    "odp-linked-image-http",
}


@dataclass
class Lab:
    root: Path
    traps: Traps
    hits: list[str] = field(default_factory=list)


class _Canary(http.server.BaseHTTPRequestHandler):
    hits: list[str]

    def parse_request(self) -> bool:
        understood = super().parse_request()
        self.server.hits.append(f"{self.command} {self.path}")  # type: ignore[attr-defined]
        return understood

    def _reply(self) -> None:
        body = f"{MARKER} from the canary\n".encode()
        self.send_response(200)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        if self.command != "HEAD":
            self.wfile.write(body)

    do_GET = do_HEAD = do_OPTIONS = do_PROPFIND = _reply

    def log_message(self, *_args: object) -> None:
        return None


@pytest.fixture(scope="module")
def lab(tmp_path_factory: pytest.TempPathFactory) -> Iterator[Lab]:
    root = tmp_path_factory.mktemp("rendition-lab")
    (root / "marker.txt").write_text(f"{MARKER} local text\n")
    (root / "marker.csv").write_text(f"{MARKER},cell\n")
    Image.new("RGB", (120, 120), MAGENTA).save(root / "marker.png")
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Canary)
    server.hits = []  # type: ignore[attr-defined]
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield Lab(
        root=root,
        traps=Traps(
            text_url=(root / "marker.txt").as_uri(),
            image_url=(root / "marker.png").as_uri(),
            csv_url=(root / "marker.csv").as_uri(),
            canary_url=f"http://127.0.0.1:{server.server_address[1]}",
            macro_proof=root / "macro_ran.txt",
        ),
        hits=server.hits,  # type: ignore[attr-defined]
    )
    server.shutdown()
    thread.join(timeout=5)


def _is_magenta(colour: tuple[int, int, int]) -> bool:
    """Close, not exact: scaling a picture onto a page shifts a channel by a step."""
    return all(abs(channel - wanted) <= 8 for channel, wanted in zip(colour, MAGENTA, strict=True))


def _escaped(pdf: Path) -> list[str]:
    """What reached the page: the marker's text, or the magenta image."""
    found: list[str] = []
    text = " ".join(page.extract_text() or "" for page in pypdf.PdfReader(pdf).pages)
    if MARKER in text:
        found.append("text")
    document = pypdfium2.PdfDocument(pdf)
    try:
        for index in range(len(document)):
            colours = document[index].render(scale=0.5).to_pil().convert("RGB").getcolors(1 << 22)
            if any(count > 20 and _is_magenta(colour) for count, colour in colours or []):
                found.append("image")
                break
    finally:
        document.close()
    return found


def _written(vector: Vector, lab: Lab, work: Path) -> Path:
    path = work / f"crafted.{vector.extension}"
    path.write_bytes(vector.build(lab.traps))
    return path


def _reset(lab: Lab) -> None:
    lab.hits.clear()
    lab.traps.macro_proof.unlink(missing_ok=True)


def _convert_stock(vector: Vector, lab: Lab, work: Path) -> Path | None:
    """The same file through a stock profile with no input hardening."""
    source = _written(vector, lab, work)
    write_profile(work / "profile", "<oor:items xmlns:oor='http://openoffice.org/2001/registry'/>")
    office_format = sniff_office_format(source)
    command = soffice_command(SOFFICE, source, office_format, work, OPTIONS)
    (work / "out").mkdir()
    # Not ``run_soffice``: its environment closes the network by itself.
    subprocess.run(command, cwd=work, env={"HOME": str(work)}, timeout=120, capture_output=True)
    produced = work / "out" / "crafted.pdf"
    return produced if produced.is_file() else None


@pytest.mark.parametrize("vector", VECTORS, ids=lambda vector: vector.name)
def test_hostile_document_cannot_escape(vector: Vector, lab: Lab, tmp_path: Path) -> None:
    _reset(lab)
    source = _written(vector, lab, tmp_path)
    office_format = sniff_office_format(source)
    assert office_format is not None, "the crafted file must pass the sniffer to prove anything"
    work = tmp_path / "work"
    work.mkdir()

    try:
        produced = convert_to_pdf(source, office_format, work, soffice=SOFFICE, options=OPTIONS)
    except ConversionError as refused:
        assert "out of time" not in str(refused), "a timeout proves nothing"
        produced = None

    assert lab.hits == [], "the document reached the network"
    assert not lab.traps.macro_proof.exists(), "a macro ran"
    assert (produced is None) == (vector.name in REFUSED)
    if produced is not None:
        assert _escaped(produced) == []
        text = " ".join(page.extract_text() or "" for page in pypdf.PdfReader(produced).pages)
        assert FIRST_PAGE_TEXT.split()[0] in text, "the rendition must still show the document"


# Vectors a stock LibreOffice 26.2 falls for, found by running every vector above
# through ``_convert_stock``. Each is held shut by the production path.
LIVE_ON_STOCK_PROFILE = {
    "docx-linked-image-file": "image",
    "docx-svg-image-file": "image",
    "docx-svg-disguised-file": "image",
    "docx-svg-gzip-file": "image",
    "docx-svg-nested-data-file": "image",
    "docx-html-chunk-file": "image",
    "odt-linked-image-file": "image",
    "odt-svg-image-file": "image",
    "odp-linked-image-file": "image",
    "xlsx-linked-image-file": "image",
    "pptx-linked-image-file": "image",
    "pptx-svg-image-file": "image",
    "pptx-ole-link-file": "text",
    "rtf-includepicture-file": "image",
    "rtf-pict-svg-file": "image",
}


# Measured on Debian's 25.2.3: these two do nothing there even on a stock profile,
# because it imports no HTML chunk and does not follow a PPTX OLE link. Each build
# has a fixed list, so a control can neither fail quietly nor go unnoticed if a new
# LibreOffice brings a trap back to life.
INERT_ON_THE_IMAGE = {"docx-html-chunk-file", "pptx-ole-link-file"}
LIVE_HERE = {
    name: kind
    for name, kind in LIVE_ON_STOCK_PROFILE.items()
    if not (IN_IMAGE and name in INERT_ON_THE_IMAGE)
}
LIVE_VECTORS = [vector for vector in VECTORS if vector.name in LIVE_HERE]


@pytest.mark.parametrize("vector", LIVE_VECTORS, ids=lambda vector: vector.name)
def test_control_the_trap_is_live_on_a_stock_profile(vector: Vector, lab: Lab, tmp_path: Path) -> None:
    _reset(lab)
    produced = _convert_stock(vector, lab, tmp_path)
    assert produced is not None
    assert LIVE_HERE[vector.name] in _escaped(produced)


@pytest.mark.skipif(not IN_IMAGE, reason="these two are live on 26.2")
@pytest.mark.parametrize(
    "vector",
    [vector for vector in VECTORS if vector.name in INERT_ON_THE_IMAGE],
    ids=lambda vector: vector.name,
)
def test_control_what_the_image_build_does_not_do_stays_that_way(
    vector: Vector, lab: Lab, tmp_path: Path
) -> None:
    _reset(lab)
    produced = _convert_stock(vector, lab, tmp_path)
    assert produced is not None
    assert _escaped(produced) == []


def _stock_run(input_path: Path, work: Path) -> Path | None:
    """``input_path`` on a stock profile, with no input hardening of its own."""
    write_profile(work / "profile", "<oor:items xmlns:oor='http://openoffice.org/2001/registry'/>")
    (work / "out").mkdir()
    command = soffice_command(SOFFICE, input_path, sniff_office_format(input_path), work, OPTIONS)
    subprocess.run(command, cwd=work, env={"HOME": str(work)}, timeout=240, capture_output=True)
    produced = work / "out" / f"{input_path.stem}.pdf"
    return produced if produced.is_file() else None


def _on_production_profile(source: Path, work: Path) -> Path:
    """The raw file on the profile production would use for its format."""
    office_format = sniff_office_format(source)
    write_profile(work / "profile", profile_settings_xml())
    (work / "out").mkdir()
    command = soffice_command(SOFFICE, source, office_format, work, OPTIONS)
    subprocess.run(command, cwd=work, env={"HOME": str(work)}, timeout=240, capture_output=True)
    return work / "out" / f"{source.stem}.pdf"


# An SVG hidden in an RTF picture needs a real RTF parser to find. The profile removes
# the SVG importer outright, and the OS wall stands behind that.
NOT_THE_COPYS_JOB = {"rtf-pict-svg-file"}


@pytest.mark.parametrize(
    "vector",
    [vector for vector in LIVE_VECTORS if vector.name not in NOT_THE_COPYS_JOB],
    ids=lambda vector: vector.name,
)
def test_the_rebuilt_copy_stops_it_with_no_help_from_the_profile(
    vector: Vector, lab: Lab, tmp_path: Path
) -> None:
    _reset(lab)
    source = _written(vector, lab, tmp_path)
    office_format = sniff_office_format(source)
    copy = tmp_path / f"copy.{office_format.extension}"
    try:
        prepare_input(source, office_format, copy)
    except (ConversionError, ValueError):
        return
    work = tmp_path / "work"
    work.mkdir()
    produced = _stock_run(copy, work)
    assert lab.hits == []
    assert produced is None or _escaped(produced) == []


# The RTF importer opens this path itself, past every setting: the copy and the OS
# wall are what stop it.
NOT_THE_PROFILES_JOB = {"rtf-includepicture-file"}


@pytest.mark.parametrize(
    "vector",
    [vector for vector in LIVE_VECTORS if vector.name not in NOT_THE_PROFILES_JOB],
    ids=lambda vector: vector.name,
)
def test_the_profile_stops_it_with_no_help_from_the_copy(
    vector: Vector, lab: Lab, tmp_path: Path
) -> None:
    _reset(lab)
    source = _written(vector, lab, tmp_path)
    work = tmp_path / "work"
    work.mkdir()
    produced = _on_production_profile(source, work)
    assert lab.hits == []
    assert not produced.is_file() or _escaped(produced) == []


def test_control_the_canary_sees_a_stock_profile_reach_the_network(lab: Lab, tmp_path: Path) -> None:
    _reset(lab)
    vector = next(item for item in VECTORS if item.name == "pptx-ole-link-http")
    _convert_stock(vector, lab, tmp_path)
    assert any("pptx-ole.bin" in hit for hit in lab.hits)


# ------------------------------------------------------------ ordinary documents


def _office_pdf(content: bytes, work: Path, name: str) -> Path:
    source = work / name
    source.write_bytes(content)
    target = work / "work"
    target.mkdir()
    return convert_to_pdf(
        source, sniff_office_format(source), target, soffice=SOFFICE, options=OPTIONS
    )


def _page_texts(pdf: Path) -> list[str]:
    return [page.extract_text() or "" for page in pypdf.PdfReader(pdf).pages]


def test_hidden_slides_stay_so_a_slide_number_is_a_page_number(tmp_path: Path) -> None:
    deck = Presentation()
    for number in range(1, 5):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Slide marker {number}"
    deck.slides[1]._element.set("show", "0")
    buffer = io.BytesIO()
    deck.save(buffer)

    pages = _page_texts(_office_pdf(buffer.getvalue(), tmp_path, "deck.pptx"))

    assert len(pages) == 4
    for number, page in enumerate(pages, start=1):
        assert f"Slide marker {number}" in page


def test_wide_sheet_fits_the_page_and_each_page_names_its_sheet(tmp_path: Path) -> None:
    workbook = Workbook()
    wide = workbook.active
    wide.title = "Wide Ledger"
    wide.append([f"Column {index}" for index in range(1, 31)])
    workbook.create_sheet("Second Tab").append(["Short"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    pages = _page_texts(_office_pdf(buffer.getvalue(), tmp_path, "book.xlsx"))

    assert len(pages) == 2
    assert "Wide Ledger" in pages[0] and "Column 30" in pages[0]
    assert "Second Tab" in pages[1]


def test_ordinary_picture_survives_the_rebuild(tmp_path: Path) -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (120, 120), MAGENTA).save(buffer, "PNG")
    picture = (
        '<w:p><w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="2000000" cy="2000000"/><wp:docPr id="1" name="p"/><a:graphic>'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic><pic:nvPicPr><pic:cNvPr id="2" name="p"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="rIdPicture"/><a:stretch><a:fillRect/></a:stretch>'
        '</pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2000000" cy="2000000"/>'
        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )
    content = crafted_docx(
        picture,
        rels=[("rIdPicture", f"{OFFICE_RELS}/image", "media/picture.png", False)],
        extra_parts={"word/media/picture.png": buffer.getvalue()},
        extra_types={"/word/media/picture.png": "image/png"},
    )
    assert "image" in _escaped(_office_pdf(content, tmp_path, "pictured.docx"))


def test_chart_and_picture_survive_in_a_real_deck(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["Northern parcel", "Southern parcel"]
    data.add_series("Acres", (12, 9))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data
    )
    picture = io.BytesIO()
    Image.new("RGB", (120, 120), MAGENTA).save(picture, "PNG")
    picture.seek(0)
    slide.shapes.add_picture(picture, Inches(7), Inches(2), Inches(2), Inches(2))
    buffer = io.BytesIO()
    deck.save(buffer)

    produced = _office_pdf(buffer.getvalue(), tmp_path, "deck.pptx")

    assert "Northern parcel" in " ".join(_page_texts(produced))
    assert "image" in _escaped(produced)


@pytest.mark.parametrize(
    ("extension", "filter_name", "words"),
    [
        ("odt", "writer8", "Rebuilt opendocument words"),
        ("ods", "calc8", "Rebuilt opendocument words"),
        ("odp", "impress8", "Rebuilt opendocument words"),
    ],
)
def test_opendocument_files_saved_by_libreoffice_survive_the_rebuild(
    tmp_path: Path, extension: str, filter_name: str, words: str
) -> None:
    """Real packages, with the manifest, settings, and metadata LibreOffice writes."""
    made_by = {"odt": "docx", "ods": "xlsx", "odp": "pptx"}[extension]
    if made_by == "docx":
        seed = crafted_docx(f"<w:p><w:r><w:t>{words}</w:t></w:r></w:p>")
    elif made_by == "xlsx":
        workbook = Workbook()
        workbook.active.append([words])
        buffer = io.BytesIO()
        workbook.save(buffer)
        seed = buffer.getvalue()
    else:
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = words
        buffer = io.BytesIO()
        deck.save(buffer)
        seed = buffer.getvalue()
    source = tmp_path / f"seed.{made_by}"
    source.write_bytes(seed)
    subprocess.run(
        [
            SOFFICE,
            f"-env:UserInstallation={(tmp_path / 'maker').as_uri()}",
            "--headless",
            "--convert-to",
            extension,
            "--outdir",
            str(tmp_path),
            str(source),
        ],
        env={"HOME": str(tmp_path)},
        timeout=240,
        capture_output=True,
    )
    saved = tmp_path / f"seed.{extension}"
    assert sniff_office_format(saved).input_filter == filter_name
    work = tmp_path / "work"
    work.mkdir()
    produced = convert_to_pdf(saved, sniff_office_format(saved), work, soffice=SOFFICE, options=OPTIONS)
    assert words in " ".join(_page_texts(produced))


def test_rtf_field_shows_its_saved_result_without_the_instruction(tmp_path: Path) -> None:
    content = (
        b"{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\f0 Plain words\\par "
        b"{\\field{\\*\\fldinst DATE}{\\fldrslt Saved result}}\\par}"
    )
    text = " ".join(_page_texts(_office_pdf(content, tmp_path, "note.rtf")))
    assert "Plain words" in text
    assert "Saved result" in text


def test_word_passage_is_found_on_its_page_through_the_real_route(tmp_path: Path) -> None:
    get_store.cache_clear()
    store = get_store()
    store.knowledge_originals = KnowledgeOriginalStore(tmp_path / "originals")
    config_id = "knowledge-rendition-route"
    store.knowledge_configs[config_id] = KnowledgeConfig(
        id=config_id,
        tenant_id="tenant-example",
        name="Rendition route",
        source_type="upload",
        enabled=True,
        acl_group_ids=["group-litigation"],
    )
    passage = "The Thornfield Orchard access covenant opens the eastern gate at harvest."
    filler = " ".join(f"Recorded boundary marker {index} of the lease." for index in range(60))
    sections = [
        f"<w:p><w:pPr><w:pageBreakBefore/></w:pPr><w:r><w:t>Section {number}. {filler}</w:t></w:r></w:p>"
        for number in range(1, 5)
    ]
    sections.append(f"<w:p><w:r><w:t>{passage}</w:t></w:r></w:p>")
    client = TestClient(app)
    response = client.post(
        f"/api/knowledge/{config_id}/documents",
        files=[("files", ("lease.docx", crafted_docx("".join(sections)), "application/octet-stream"))],
        headers={"x-aperture-user": "user-admin"},
    )
    assert response.status_code == 200
    document = response.json()["documents"][0]
    chunk = next(
        item for item in store.knowledge_chunks_for(config_id) if "Thornfield" in item.text
    )

    view = client.get(
        f"/api/knowledge/{config_id}/documents/{document['id']}/view",
        params={"chunk_id": chunk.id},
        headers={"x-aperture-user": "user-jane"},
    ).json()

    assert view["mode"] == "pdf"
    assert view["page_count"] == 5
    assert view["target"]["page"] == 5
    assert view["target"]["exact"] is True
    served = client.get(view["file_url"], headers={"x-aperture-user": "user-jane"})
    assert served.headers["content-type"] == "application/pdf"
    assert served.content.startswith(b"%PDF-")
    get_store.cache_clear()


# ------------------------------------------------- reading a file through a document


WALLED_PHRASE = "WALLEDPHRASE5519 confidential settlement figure"


def _text_pdf(text: str) -> bytes:
    document = pikepdf.new()
    font = pikepdf.Dictionary(
        Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1, BaseFont=pikepdf.Name.Helvetica
    )
    page = document.add_blank_page(page_size=(612, 792))
    page.Resources = pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font))
    # The magenta block is what the check looks for: an embedded page can lose its
    # text encoding on the way through, but it cannot lose its picture.
    page.Contents = document.make_stream(
        f"1 0 1 rg 72 200 468 400 re f 0 0 0 rg BT /F1 12 Tf 72 700 Td ({text}) Tj ET".encode()
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _shows_walled_text(pdf: Path) -> bool:
    return "WALLEDPHRASE5519" in " ".join(_page_texts(pdf)) or "image" in _escaped(pdf)


def _stock_conversion(content: bytes, work: Path, extension: str = "docx") -> Path | None:
    """The file as uploaded, on a stock profile, with every input layer skipped."""
    source = work / f"crafted.{extension}"
    source.write_bytes(content)
    write_profile(work / "profile", "<oor:items xmlns:oor='http://openoffice.org/2001/registry'/>")
    (work / "out").mkdir()
    command = soffice_command(SOFFICE, source, sniff_office_format(source), work, OPTIONS)
    subprocess.run(command, cwd=work, env={"HOME": str(work)}, timeout=240, capture_output=True)
    produced = work / "out" / "crafted.pdf"
    return produced if produced.is_file() else None


_ALTERNATE_CONTENT = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
    '<mc:Choice Requires="w">{}</mc:Choice></mc:AlternateContent>'
)
# How a file is pulled onto a page by path. Each takes the file's URL.
WALL_CROSSINGS: dict[str, Callable[[str], bytes]] = {
    "linked-picture": linked_picture,
    "masked-nested-document": lambda url: masked_nested_document(linked_picture(url)),
    "retyped": retyped_chunk,
    "unknown-type": lambda url: retyped_chunk(url, kind="urn:not-a-chunk"),
    "utf-16": lambda url: retyped_chunk(url, document_encoding="UTF-16"),
    "alternate-content": lambda url: retyped_chunk(url, wrap=_ALTERNATE_CONTENT),
}
# Live on a stock profile on each build. 25.2.3 imports no HTML chunk, so the HTML
# crossings prove nothing there. The control below fails if a listed one stops firing.
PICTURE_CROSSINGS = ["linked-picture", "masked-nested-document"]
LIVE_CROSSINGS = (
    PICTURE_CROSSINGS if IN_IMAGE else [*PICTURE_CROSSINGS, "retyped", "unknown-type", "utf-16"]
)


@pytest.fixture
def walled_store(tmp_path: Path) -> Iterator[object]:
    """Two tenants. The first holds a walled PDF that nobody below may read."""
    get_store.cache_clear()
    store = get_store()
    store.knowledge_originals = KnowledgeOriginalStore(tmp_path / "originals")
    store.groups["group-falcon"] = Group(
        id="group-falcon",
        tenant_id="tenant-example",
        name="Project Falcon",
        distinguished_name="CN=Project Falcon",
        entra_object_id="group-falcon",
    )
    store.tenants["tenant-rival"] = store.tenants["tenant-example"].model_copy(
        update={"id": "tenant-rival", "slug": "rival", "custom_domain": "rival.example.test"}
    )
    store.users["user-rival"] = User(
        id="user-rival",
        tenant_id="tenant-rival",
        email="rival@example.test",
        display_name="Rival admin",
        role=Role.TENANT_ADMIN,
    )
    for config_id, tenant_id in (("kb-walled", "tenant-example"), ("kb-attack", "tenant-example"), ("kb-rival", "tenant-rival")):
        store.knowledge_configs[config_id] = KnowledgeConfig(
            id=config_id,
            tenant_id=tenant_id,
            name=config_id,
            source_type="upload",
            enabled=True,
        )
    wall = KnowledgeDocumentWall(source="manual", group_ids=["group-falcon"])
    documents = [
        KnowledgeDocument(
            id=document_id,
            knowledge_config_id="kb-walled",
            tenant_id="tenant-example",
            name=name,
            source_uri=f"upload://{name}",
            source_type="upload",
            chunk_count=0,
            updated_at="now",
            wall=wall,
        )
        for document_id, name in (("doc-walled", "memo.pdf"), ("doc-walled-scan", "stamp.png"))
    ]
    store.knowledge_documents["kb-walled"] = documents
    store.vector_store.replace_config("kb-walled", documents, [])
    picture = io.BytesIO()
    Image.new("RGB", (120, 120), MAGENTA).save(picture, "PNG")
    store.knowledge_originals.save(
        DocumentRef("tenant-example", "kb-walled", "doc-walled"), _text_pdf(WALLED_PHRASE)
    )
    store.knowledge_originals.save(
        DocumentRef("tenant-example", "kb-walled", "doc-walled-scan"), picture.getvalue()
    )
    yield store
    get_store.cache_clear()


def _walled_url(store: object, crossing: str = "retyped") -> str:
    """A linked picture is read as a picture, so it aims at the walled scan."""
    document = "doc-walled-scan" if crossing in PICTURE_CROSSINGS else "doc-walled"
    ref = DocumentRef("tenant-example", "kb-walled", document)
    return store.knowledge_originals.path(ref).as_uri()


def _read_through_upload(store: object, config_id: str, user_id: str, crafted: bytes, tmp_path: Path) -> Path | None:
    """What the attacker does: upload a file, open it, keep the rendition."""
    client = TestClient(app)
    identity = {"x-aperture-user": user_id}
    response = client.post(
        f"/api/knowledge/{config_id}/documents",
        files=[("files", ("note.docx", crafted, "application/octet-stream"))],
        headers=identity,
    )
    assert response.status_code == 200
    document = response.json()["documents"][-1]
    view = client.get(
        f"/api/knowledge/{config_id}/documents/{document['id']}/view", headers=identity
    ).json()
    if view["mode"] != "pdf":
        return None
    served = client.get(view["file_url"], headers=identity)
    rendition = tmp_path / "attacker-rendition.pdf"
    rendition.write_bytes(served.content)
    return rendition


@pytest.mark.parametrize("crossing", WALL_CROSSINGS, ids=str)
def test_screened_admin_cannot_read_a_walled_pdf_through_an_upload(
    walled_store: object, crossing: str, tmp_path: Path
) -> None:
    denied = TestClient(app).get(
        "/api/knowledge/kb-walled/documents/doc-walled/view",
        headers={"x-aperture-user": "user-admin"},
    )
    assert denied.status_code == 404

    crafted = WALL_CROSSINGS[crossing](_walled_url(walled_store, crossing))
    rendition = _read_through_upload(walled_store, "kb-attack", "user-admin", crafted, tmp_path)

    assert rendition is None or not _shows_walled_text(rendition)


def test_nested_document_never_reaches_the_page_through_the_real_routes(
    walled_store: object, tmp_path: Path
) -> None:
    crafted = inert_nested_document()
    rendition = _read_through_upload(walled_store, "kb-attack", "user-admin", crafted, tmp_path)
    assert rendition is None or NESTED_WORDS not in " ".join(_page_texts(rendition))


def test_control_a_nested_document_reaches_the_page_on_a_stock_profile(tmp_path: Path) -> None:
    produced = _stock_conversion(inert_nested_document(), tmp_path)
    assert produced is not None
    assert NESTED_WORDS in " ".join(_page_texts(produced))


def test_masked_parts_do_not_survive_into_the_copy(tmp_path: Path) -> None:
    source = tmp_path / "crafted.docx"
    source.write_bytes(inert_nested_document())
    copy = tmp_path / "copy.docx"
    prepare_input(source, sniff_office_format(source), copy)
    with zipfile.ZipFile(copy) as archive:
        assert not [name for name in archive.namelist() if name.endswith(".emf")]


def _deck_with_a_chart_and_an_ole_link(file_url: str) -> bytes:
    """A real deck, a real chart, and an OLE link labelled as a hyperlink."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    data = CategoryChartData()
    data.categories = ["Northern parcel"]
    data.add_series("Acres", (12,))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), data
    )
    relationship = slide.part.relate_to(file_url, f"{OFFICE_RELS}/hyperlink", is_external=True)
    frame = ole_frame().replace("rIdTrap", relationship).replace(
        "<p:graphicFrame>",
        f'<p:graphicFrame xmlns:p="{SLIDE}" xmlns:a="{DRAWING}" xmlns:r="{OFFICE_RELS}">',
        1,
    )
    slide.shapes._spTree.append(parse_xml(frame))
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_control_raw_ole_link_reads_a_local_file_once_active_content_is_on(
    lab: Lab, tmp_path: Path
) -> None:
    """Why active content stays off unless a document has nothing active in it. 26.2
    follows the link; 25.2.3 does not, and the proof fails if that changes."""
    source = tmp_path / "crafted.pptx"
    source.write_bytes(_deck_with_a_chart_and_an_ole_link(lab.traps.text_url))
    work = tmp_path / "work"
    work.mkdir()
    write_profile(work / "profile", profile_settings_xml(draws_embedded_charts=True))
    (work / "out").mkdir()
    command = soffice_command(SOFFICE, source, sniff_office_format(source), work, OPTIONS)
    subprocess.run(command, cwd=work, env={"HOME": str(work)}, timeout=240, capture_output=True)
    assert ("text" in _escaped(work / "out" / "crafted.pdf")) is not IN_IMAGE


def test_raw_ole_link_reads_nothing_on_the_profile_every_other_document_gets(
    lab: Lab, tmp_path: Path
) -> None:
    source = tmp_path / "crafted.pptx"
    source.write_bytes(_deck_with_a_chart_and_an_ole_link(lab.traps.text_url))
    work = tmp_path / "work"
    work.mkdir()
    produced = _on_production_profile(source, work)
    assert produced.is_file()
    assert _escaped(produced) == []


def test_chart_beside_an_ole_link_does_not_turn_active_content_on(
    lab: Lab, tmp_path: Path
) -> None:
    source = tmp_path / "crafted.pptx"
    source.write_bytes(_deck_with_a_chart_and_an_ole_link(lab.traps.text_url))
    copy = tmp_path / "copy.pptx"
    assert prepare_input(source, sniff_office_format(source), copy) is False
    work = tmp_path / "work"
    work.mkdir()
    produced = convert_to_pdf(
        source, sniff_office_format(source), work, soffice=SOFFICE, options=OPTIONS
    )
    assert _escaped(produced) == []


def test_profile_alone_keeps_the_html_import_shut(walled_store: object, tmp_path: Path) -> None:
    source = tmp_path / "crafted.docx"
    source.write_bytes(retyped_chunk(_walled_url(walled_store)))
    work = tmp_path / "work"
    work.mkdir()
    produced = _on_production_profile(source, work)
    assert not produced.is_file() or not _shows_walled_text(produced)


@pytest.mark.parametrize("crossing", WALL_CROSSINGS, ids=str)
def test_other_tenant_cannot_read_a_file_through_an_upload(
    walled_store: object, crossing: str, tmp_path: Path
) -> None:
    denied = TestClient(app).get(
        "/api/knowledge/kb-walled/documents/doc-walled/original",
        headers={"x-aperture-user": "user-rival"},
    )
    assert denied.status_code == 404

    crafted = WALL_CROSSINGS[crossing](_walled_url(walled_store, crossing))
    rendition = _read_through_upload(walled_store, "kb-rival", "user-rival", crafted, tmp_path)

    assert rendition is None or not _shows_walled_text(rendition)


@pytest.mark.parametrize("crossing", LIVE_CROSSINGS, ids=str)
def test_control_a_stock_profile_shows_the_walled_text(
    walled_store: object, crossing: str, tmp_path: Path
) -> None:
    """The same bytes do read the file with the defences out of the way, so the
    two tests above would fail if this class of read came back."""
    crafted = WALL_CROSSINGS[crossing](_walled_url(walled_store, crossing))
    produced = _stock_conversion(crafted, tmp_path)
    assert produced is not None
    assert _shows_walled_text(produced)
