"""Text that plain body extraction used to drop, and AI page reading.

Covers Word headers/footers/footnotes/comments, slide speaker notes and
pictures, PDF reviewer comments and filled form fields, email attachments, and
the vision page-reader hook used for handwriting, stamps, and images. All files
are synthetic and built in the test.
"""

from __future__ import annotations

import zipfile
from email.message import EmailMessage
from io import BytesIO
from typing import Any

import pytest
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    FloatObject,
    NameObject,
    NumberObject,
    TextStringObject,
)

from app.core import knowledge_ingestion
from app.core.knowledge_ingestion import extract_document, extract_segments

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
TYPED = "This typed clause is comfortably longer than forty characters of text."


class FakeReader:
    """A vision page reader that records what it was asked to read."""

    concurrency = 2

    def __init__(self, answers: dict[str, str | None] | None = None) -> None:
        self.answers = answers or {}
        self.calls: list[tuple[str, str]] = []

    def read_page(self, image: Any, *, mode: str, label: str) -> str | None:
        assert image.width > 0 and image.height > 0
        self.calls.append((mode, label))
        return self.answers.get(mode)


def _para(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def _docx(extra: dict[str, str], media: dict[str, bytes] | None = None) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "word/document.xml",
            f"<w:document {_W}><w:body>{_para('Main body of the agreement.')}</w:body></w:document>",
        )
        for name, content in extra.items():
            archive.writestr(name, content)
        for name, content in (media or {}).items():
            archive.writestr(name, content)
    return buffer.getvalue()


def _png(size: tuple[int, int] = (400, 300), *, noise: bool = True) -> bytes:
    image = Image.new("RGB", size, "white")
    if noise:
        # Incompressible pixels keep the encoded image above the "decoration" floor.
        image = Image.effect_noise(size, 80).convert("RGB")
    draw = ImageDraw.Draw(image)
    draw.text((10, 10), "Signed: J. Sample", fill="black")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _by_locator(segments: list[Any]) -> dict[str | None, str]:
    return {segment.locator: segment.text for segment in segments}


# ------------------------------------------------------------------- Word


def test_word_headers_footers_footnotes_endnotes_and_comments_are_indexed() -> None:
    content = _docx(
        {
            "word/header1.xml": f"<w:hdr {_W}>{_para('PRIVILEGED AND CONFIDENTIAL')}</w:hdr>",
            "word/header2.xml": f"<w:hdr {_W}>{_para('PRIVILEGED AND CONFIDENTIAL')}</w:hdr>",
            "word/footer1.xml": f"<w:ftr {_W}>{_para('Draft 3 - Example Partners')}</w:ftr>",
            "word/footnotes.xml": (
                f"<w:footnotes {_W}><w:footnote w:id='1'>{_para('See Exhibit B.')}</w:footnote></w:footnotes>"
            ),
            "word/endnotes.xml": (
                f"<w:endnotes {_W}><w:endnote w:id='1'>{_para('Final endnote.')}</w:endnote></w:endnotes>"
            ),
            "word/comments.xml": (
                f"<w:comments {_W}><w:comment w:id='0' w:author='Avery Example' w:date='2026-03-01T10:00:00Z'>"
                f"{_para('Confirm the cap with the client.')}</w:comment></w:comments>"
            ),
        }
    )
    segments = _by_locator(extract_segments("agreement.docx", content))
    assert segments[None] == "Main body of the agreement."
    assert segments["Headers and footers"] == "PRIVILEGED AND CONFIDENTIAL\n\nDraft 3 - Example Partners"
    assert segments["Footnotes"] == "See Exhibit B."
    assert segments["Endnotes"] == "Final endnote."
    assert segments["Comments"] == "Avery Example: Confirm the cap with the client."


def test_word_embedded_pictures_are_read_by_the_vision_reader() -> None:
    reader = FakeReader({"image": "Handwritten note: approve at 42M."})
    content = _docx(
        {},
        media={
            "word/media/image1.png": _png(),
            "word/media/logo.png": _png((20, 20), noise=False),
        },
    )
    segments = _by_locator(extract_segments("memo.docx", content, vision=reader))
    assert segments["Embedded image 1"] == "Handwritten note: approve at 42M."
    # The tiny logo is decoration and is never sent to the model.
    assert reader.calls == [("image", "Embedded image 1")]


def test_word_pictures_are_skipped_without_a_vision_reader() -> None:
    content = _docx({}, media={"word/media/image1.png": _png()})
    assert list(_by_locator(extract_segments("memo.docx", content))) == [None]


# ---------------------------------------------------------------- PowerPoint


def test_slide_speaker_notes_and_pictures_are_indexed() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Q3 litigation budget"
    slide.notes_slide.notes_text_frame.text = "Mention the reserve increase."
    slide.shapes.add_picture(BytesIO(_png()), Inches(1), Inches(2))
    buffer = BytesIO()
    deck.save(buffer)

    reader = FakeReader({"image": "Chart: spend rises 12% in Q3."})
    segments = _by_locator(extract_segments("budget.pptx", buffer.getvalue(), vision=reader))
    assert segments["Slide 1"] == (
        "Q3 litigation budget\n\nImage: Chart: spend rises 12% in Q3.\n\n"
        "Speaker notes: Mention the reserve increase."
    )


# ----------------------------------------------------------------------- PDF


def _pdf(
    *,
    text: str | None = None,
    image: bool = False,
    annotations: list[DictionaryObject] | None = None,
    fields: list[DictionaryObject] | None = None,
) -> bytes:
    if image:
        buffer = BytesIO()
        Image.open(BytesIO(_png((600, 800)))).save(buffer, format="PDF")
        writer = PdfWriter(clone_from=PdfReader(BytesIO(buffer.getvalue())))
    else:
        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
    page = writer.pages[0]
    if text is not None:
        _add_text(writer, page, text)
    refs = []
    for annotation in [*(annotations or []), *(fields or [])]:
        refs.append(writer._add_object(annotation))
    if refs:
        page[NameObject("/Annots")] = ArrayObject(refs)
    if fields:
        writer._root_object[NameObject("/AcroForm")] = DictionaryObject(
            {NameObject("/Fields"): ArrayObject(refs[len(annotations or []) :])}
        )
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def _add_text(writer: PdfWriter, page: Any, text: str) -> None:
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = page.get("/Resources")
    resources = resources.get_object() if resources is not None else DictionaryObject()
    resources[NameObject("/Font")] = DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    page[NameObject("/Resources")] = resources
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 700 Td ({text}) Tj ET".encode("latin-1"))
    existing = page.get("/Contents")
    contents = ArrayObject([existing] if existing is not None else [])
    contents.append(writer._add_object(stream))
    page[NameObject("/Contents")] = contents


def _annotation(subtype: str, contents: str | None = None, author: str | None = None) -> DictionaryObject:
    annotation = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject(subtype),
            NameObject("/Rect"): ArrayObject([FloatObject(50), FloatObject(50), FloatObject(200), FloatObject(120)]),
        }
    )
    if contents is not None:
        annotation[NameObject("/Contents")] = TextStringObject(contents)
    if author is not None:
        annotation[NameObject("/T")] = TextStringObject(author)
    return annotation


def _text_field(name: str, label: str | None, value: str) -> DictionaryObject:
    field = _annotation("/Widget")
    field[NameObject("/FT")] = NameObject("/Tx")
    field[NameObject("/T")] = TextStringObject(name)
    field[NameObject("/V")] = TextStringObject(value)
    field[NameObject("/F")] = NumberObject(4)
    if label is not None:
        field[NameObject("/TU")] = TextStringObject(label)
    return field


def test_pdf_reviewer_comments_are_kept_with_their_page() -> None:
    content = _pdf(
        text=TYPED,
        annotations=[
            _annotation("/Text", "Check this date against the docket.", "Jordan Sample"),
            _annotation("/Link"),
        ],
    )
    segments = _by_locator(extract_segments("brief.pdf", content, ocr_enabled=False))
    assert segments["Page 1"] == (
        f"{TYPED}\n\nComments on this page:\n\n- Jordan Sample: Check this date against the docket."
    )


def test_pdf_filled_form_field_values_are_indexed() -> None:
    content = _pdf(
        text=TYPED,
        fields=[
            _text_field("topmostSubform[0].Page1[0].f1_01[0]", "Client name", "Kestrel Holdings"),
            _text_field("matter_no", None, "20419"),
            _text_field("blank", "Unused", ""),
        ],
    )
    segments = _by_locator(extract_segments("intake.pdf", content, ocr_enabled=False))
    assert segments["Form fields"] == "Client name: Kestrel Holdings\n\nmatter_no: 20419"


def test_scanned_pdf_page_is_read_by_the_vision_reader_first(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(knowledge_ingestion, "_ocr_image", lambda *_a, **_k: "tesseract text")
    reader = FakeReader({"full": "Handwritten: Call Jordan re: settlement by Friday."})
    content = _pdf(image=True)
    segments = _by_locator(extract_segments("scan.pdf", content, vision=reader))
    assert segments["Page 1"] == "Handwritten: Call Jordan re: settlement by Friday."
    assert reader.calls == [("full", "Page 1")]


def test_scanned_pdf_page_falls_back_to_tesseract_when_vision_fails(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(knowledge_ingestion, "_ocr_image", lambda *_a, **_k: "tesseract text")
    reader = FakeReader({"full": None})
    segments = _by_locator(extract_segments("scan.pdf", _pdf(image=True), vision=reader))
    assert segments["Page 1"] == "tesseract text"


def test_typed_page_with_a_large_image_gets_handwriting_and_image_notes() -> None:
    reader = FakeReader({"extras": "Margin note in blue ink: 'agreed - JS'."})
    segments = _by_locator(
        extract_segments("signed.pdf", _pdf(text=TYPED, image=True), vision=reader)
    )
    assert segments["Page 1"] == (
        f"{TYPED}\n\nHandwriting, stamps, and images on this page:\n\n"
        "Margin note in blue ink: 'agreed - JS'."
    )
    assert reader.calls == [("extras", "Page 1")]


def test_typed_page_with_ink_annotation_is_read_for_markings() -> None:
    reader = FakeReader({"extras": "Ink initials: JS."})
    content = _pdf(text=TYPED, annotations=[_annotation("/Ink")])
    segments = _by_locator(extract_segments("initialed.pdf", content, vision=reader))
    assert segments["Page 1"].endswith("Handwriting, stamps, and images on this page:\n\nInk initials: JS.")


def test_typed_page_without_images_is_not_sent_to_the_vision_reader() -> None:
    reader = FakeReader({"extras": "should not be used"})
    segments = _by_locator(extract_segments("plain.pdf", _pdf(text=TYPED), vision=reader))
    assert segments["Page 1"] == TYPED
    assert reader.calls == []


def test_vision_none_answer_adds_nothing() -> None:
    reader = FakeReader({"extras": None})
    segments = _by_locator(
        extract_segments("signed.pdf", _pdf(text=TYPED, image=True), vision=reader)
    )
    assert segments["Page 1"] == TYPED


# --------------------------------------------------------------------- images


def test_image_upload_prefers_vision_and_falls_back_to_tesseract(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(knowledge_ingestion, "_ocr_image", lambda *_a, **_k: "tesseract text")
    vision = FakeReader({"full": "Whiteboard: Phase 2 starts May 4."})
    assert extract_segments("board.png", _png(), vision=vision)[0].text == (
        "Whiteboard: Phase 2 starts May 4."
    )
    failing = FakeReader({"full": None})
    assert extract_segments("board.png", _png(), vision=failing)[0].text == "tesseract text"


def test_image_upload_with_ocr_off_still_uses_vision() -> None:
    vision = FakeReader({"full": "Sticky note: renew lease."})
    segments = extract_segments("note.jpg", _png(), ocr_enabled=False, vision=vision)
    assert segments[0].text == "Sticky note: renew lease."


# ---------------------------------------------------------------------- email


def test_email_attachments_are_extracted_with_their_names() -> None:
    message = EmailMessage()
    message["From"] = "avery@example.com"
    message["Subject"] = "Closing set"
    message.set_content("Attached are the closing documents.")
    message.add_attachment(
        b"Item 1: wire instructions confirmed.",
        maintype="text",
        subtype="plain",
        filename="checklist.txt",
    )
    message.add_attachment(
        _docx({}), maintype="application", subtype="octet-stream", filename="agreement.docx"
    )
    document = extract_document("closing.eml", message.as_bytes())
    segments = _by_locator(document.segments)
    assert segments[None].startswith("From: avery@example.com")
    assert segments["Attachment: checklist.txt"] == "Item 1: wire instructions confirmed."
    assert segments["Attachment: agreement.docx"] == "Main body of the agreement."


def test_nested_attachment_page_locators_name_the_attachment() -> None:
    message = EmailMessage()
    message["Subject"] = "Signed brief"
    message.set_content("See attached.")
    message.add_attachment(
        _pdf(text=TYPED), maintype="application", subtype="pdf", filename="brief.pdf"
    )
    segments = extract_document("brief.eml", message.as_bytes(), ocr_enabled=False).segments
    attachment = [segment for segment in segments if segment.locator and "brief.pdf" in segment.locator]
    assert attachment[0].locator == "Attachment: brief.pdf, Page 1"
    assert attachment[0].text == TYPED
    # Page numbers belong to the attachment, not the email, so none are claimed.
    assert attachment[0].page_start is None


def test_forwarded_message_attachment_is_extracted() -> None:
    inner = EmailMessage()
    inner["Subject"] = "Original terms"
    inner.set_content("Inner message: the cap is 1.2 million.")
    outer = EmailMessage()
    outer["Subject"] = "Fwd: terms"
    outer.set_content("Forwarding the original.")
    outer.add_attachment(inner, filename="original.eml")
    segments = _by_locator(extract_document("fwd.eml", outer.as_bytes()).segments)
    assert "Inner message: the cap is 1.2 million." in segments["Attachment: original.eml"]


def test_outlook_attachments_are_extracted(monkeypatch: pytest.MonkeyPatch) -> None:
    class Attachment:
        def __init__(self, name: str | None, data: bytes | None, mime: str | None) -> None:
            self.file_name, self.file_bytes, self.mime_type = name, data, mime

    class Message:
        sender = "avery@example.com"
        subject = "Closing set"
        sent_date = None
        message_headers: dict[str, str] = {}
        body = "See attached."
        html_body = None
        attachments = [
            Attachment("terms.txt", b"Escrow is 3 million.", "text/plain"),
            Attachment(None, b"orphan", "text/plain"),
            Attachment("broken.txt", None, "text/plain"),
        ]

    monkeypatch.setattr(knowledge_ingestion, "_parse_msg", lambda _source: Message())
    segments = _by_locator(extract_segments("closing.msg", b"msg-bytes"))
    assert segments[None] == "From: avery@example.com Subject: Closing set\n\nSee attached."
    assert segments["Attachment: terms.txt"] == "Escrow is 3 million."
    assert len(segments) == 2


def test_outlook_html_body_is_used_when_there_is_no_plain_body(monkeypatch: pytest.MonkeyPatch) -> None:
    class Message:
        sender = None
        subject = "HTML only"
        sent_date = None
        message_headers: dict[str, str] = {}
        body = None
        html_body = b"<p>Rich <b>body</b> text.</p>"
        attachments: list[object] = []

    monkeypatch.setattr(knowledge_ingestion, "_parse_msg", lambda _source: Message())
    [segment] = extract_segments("rich.msg", b"msg-bytes")
    assert segment.text.startswith("Subject: HTML only")
    assert "Rich body text." in " ".join(segment.text.split()).replace(" .", ".")


# ---------------------------------------------------------------- boundaries


@pytest.mark.parametrize(
    ("filename", "mime", "kind"),
    [
        ("upload", "application/pdf; name=brief", "pdf"),
        ("scan", "image/png", "image"),
        ("scan.tif", "application/octet-stream", "image"),
        ("notes", "text/markdown", "text"),
        ("data", "application/json", "text"),
        ("server.log", "application/octet-stream", "text"),
        ("blob", "application/octet-stream", "unknown"),
        # Only the guessed type says this is text.
        ("table.tsv", None, "text"),
    ],
)
def test_source_kind_uses_extension_or_declared_type(filename: str, mime: str | None, kind: str) -> None:
    assert knowledge_ingestion._source_kind(filename, mime) == kind


def test_headers_come_before_footers_in_numeric_order() -> None:
    content = _docx(
        {
            "word/footer1.xml": f"<w:ftr {_W}>{_para('Footer one')}</w:ftr>",
            "word/header10.xml": f"<w:hdr {_W}>{_para('Header ten')}</w:hdr>",
            "word/header2.xml": f"<w:hdr {_W}>{_para('Header two A')}{_para('Header two B')}</w:hdr>",
            "word/header.xml": f"<w:hdr {_W}>{_para('Header plain')}</w:hdr>",
        }
    )
    segments = _by_locator(extract_segments("agreement.docx", content))
    assert segments["Headers and footers"] == (
        "Header plain\n\nHeader two A\n\nHeader two B\n\nHeader ten\n\nFooter one"
    )


def _email_with(*attachments: tuple[bytes, str, str, str | None]) -> EmailMessage:
    message = EmailMessage()
    message["Subject"] = "Attachments"
    message.set_content("Body.")
    for payload, maintype, subtype, filename in attachments:
        if filename is None:
            message.add_attachment(payload, maintype=maintype, subtype=subtype)
        else:
            message.add_attachment(payload, maintype=maintype, subtype=subtype, filename=filename)
    return message


def test_unnamed_attachments_and_tiny_images_are_skipped() -> None:
    tiny = _png((16, 16), noise=False)
    message = _email_with(
        (b"no name here", "text", "plain", None),
        (tiny, "image", "png", "logo.png"),
        (b"Kept text.", "text", "plain", "kept.txt"),
    )
    vision = FakeReader({"full": "should not read the logo"})
    segments = _by_locator(extract_document("mail.eml", message.as_bytes(), vision=vision).segments)
    assert set(segments) == {None, "Attachment: kept.txt"}
    assert vision.calls == []


def test_attachment_size_limit_is_inclusive(monkeypatch: pytest.MonkeyPatch) -> None:
    payload = b"Exactly at the limit."
    message = _email_with((payload, "text", "plain", "limit.txt"))
    monkeypatch.setattr(knowledge_ingestion, "MAX_ATTACHMENT_BYTES", len(payload))
    segments = _by_locator(extract_document("mail.eml", message.as_bytes()).segments)
    assert segments["Attachment: limit.txt"] == "Exactly at the limit."
    monkeypatch.setattr(knowledge_ingestion, "MAX_ATTACHMENT_BYTES", len(payload) - 1)
    segments = _by_locator(extract_document("mail.eml", message.as_bytes()).segments)
    assert "Attachment: limit.txt" not in segments


def test_picture_attachment_at_the_decoration_floor_is_read(monkeypatch: pytest.MonkeyPatch) -> None:
    picture = _png()
    monkeypatch.setattr(knowledge_ingestion, "MIN_VISION_IMAGE_BYTES", len(picture))
    message = _email_with((picture, "image", "png", "photo.png"))
    vision = FakeReader({"full": "A photo of a signed page."})
    segments = _by_locator(extract_document("mail.eml", message.as_bytes(), vision=vision).segments)
    assert segments["Attachment: photo.png, Page 1"] == "A photo of a signed page."


def test_attachments_are_followed_two_messages_deep_and_no_further() -> None:
    def wrap(inner: EmailMessage, name: str, text: str) -> EmailMessage:
        outer = EmailMessage()
        outer["Subject"] = name
        outer.set_content(text)
        outer.add_attachment(inner, filename=f"{name}.eml")
        return outer

    deepest = EmailMessage()
    deepest["Subject"] = "deepest"
    deepest.set_content("Level three body.")
    deepest.add_attachment(b"Level three file.", maintype="text", subtype="plain", filename="three.txt")
    level_two = wrap(deepest, "two", "Level two body.")
    level_one = wrap(level_two, "one", "Level one body.")
    text = "\n".join(segment.text for segment in extract_document("top.eml", level_one.as_bytes()).segments)
    assert "Level two body." in text
    assert "Level three body." in text
    assert "Level three file." not in text


def test_image_coverage_sums_image_areas_and_caps_at_the_page() -> None:
    class Item:
        def __init__(self, bounds: tuple[float, float, float, float]) -> None:
            self.bounds = bounds

        def get_bounds(self) -> tuple[float, float, float, float]:
            return self.bounds

    class Page:
        def __init__(self, size: tuple[float, float], items: list[Item]) -> None:
            self.size, self.items, self.closed = size, items, False

        def get_size(self) -> tuple[float, float]:
            return self.size

        def get_objects(self, **_kwargs: Any) -> list[Item]:
            return self.items

        def close(self) -> None:
            self.closed = True

    class Raw:
        FPDF_PAGEOBJ_IMAGE = 3

    quarter = Page((200.0, 100.0), [Item((50, 0, 150, 50)), Item((10, 60, 10, 90))])
    assert knowledge_ingestion._image_coverage(quarter, Raw) == pytest.approx(0.25)
    assert quarter.closed
    overlapping = Page((100.0, 100.0), [Item((0, 0, 100, 100)), Item((0, 0, 50, 100))])
    assert knowledge_ingestion._image_coverage(overlapping, Raw) == 1.0
    assert knowledge_ingestion._image_coverage(Page((0.0, 100.0), [Item((0, 0, 1, 1))]), Raw) == 0.0


def test_image_heavy_pages_include_the_threshold_and_skip_missing_pages(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    coverage = {0: knowledge_ingestion.MIN_IMAGE_COVERAGE, 1: knowledge_ingestion.MIN_IMAGE_COVERAGE - 0.01}
    pages = iter([0, 1])
    monkeypatch.setattr(knowledge_ingestion, "_image_coverage", lambda page, _raw: coverage[next(pages)])
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    # Index 2 equals the page count, so it is skipped without stopping the scan.
    assert knowledge_ingestion._image_heavy_pages(buffer.getvalue(), [2, 0, 1]) == [0]


def test_markings_pass_never_falls_back_to_tesseract(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(knowledge_ingestion, "_ocr_image", lambda *_a, **_k: "tesseract text")
    reader = FakeReader({"extras": None})
    segments = _by_locator(extract_segments("signed.pdf", _pdf(text=TYPED, image=True), vision=reader))
    assert segments["Page 1"] == TYPED


def test_page_with_exactly_the_minimum_text_counts_as_typed(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(knowledge_ingestion, "_ocr_image", lambda *_a, **_k: "tesseract text")
    exact = "x" * knowledge_ingestion.MIN_PDF_PAGE_TEXT_CHARS
    reader = FakeReader({"full": "scanned reading", "extras": "markings"})
    segments = _by_locator(extract_segments("edge.pdf", _pdf(text=exact, image=True), vision=reader))
    assert segments["Page 1"] == f"{exact}\n\nHandwriting, stamps, and images on this page:\n\nmarkings"
    assert reader.calls == [("extras", "Page 1")]


def test_animated_images_read_each_frame_up_to_the_page_limit() -> None:
    frames = [Image.new("RGB", (60, 40), color) for color in ("white", "gray", "black")]
    buffer = BytesIO()
    frames[0].save(buffer, format="GIF", save_all=True, append_images=frames[1:])
    reader = FakeReader({"full": "frame text"})
    segments = extract_segments("scan.gif", buffer.getvalue(), ocr_max_pages=2, vision=reader)
    assert [(segment.locator, segment.page_start) for segment in segments] == [("Page 1", 1), ("Page 2", 2)]
    assert sorted(label for _mode, label in reader.calls) == ["Page 1", "Page 2"]


def test_slide_picture_size_floor_is_inclusive(monkeypatch: pytest.MonkeyPatch) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    picture = _png((300, 200))
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(picture), Inches(1), Inches(1))
    buffer = BytesIO()
    deck.save(buffer)
    stored = len(Presentation(BytesIO(buffer.getvalue())).slides[0].shapes[0].image.blob)
    monkeypatch.setattr(knowledge_ingestion, "MIN_VISION_IMAGE_BYTES", stored)
    reader = FakeReader({"image": "A chart."})
    segments = _by_locator(extract_segments("deck.pptx", buffer.getvalue(), vision=reader))
    assert segments["Slide 1"] == "Image: A chart."
