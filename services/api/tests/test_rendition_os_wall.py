"""The wall under every other layer, run for real inside the built API image.

As root on Linux, with the image's ``aperture-render`` account. Everywhere else
these skip, except that ``scripts/prove-rendition-image.sh`` sets
``APERTURE_RENDITION_IMAGE_PROOF=1``, which turns any such skip into a failure so
the proof can never pass by proving nothing.

The heart of it is ``test_the_wall_alone_stops_a_read_every_other_layer_missed``:
a hostile file, unsanitized, on a stock profile, aimed at a synthetic secret in
the data directory. As root it reads the secret, which shows the trap is live.
As the account the same run finds nothing, which is what holds when a vector
nobody has thought of gets past the parsers. Nothing here may xfail: a wall test
that cannot prove its point fails.
"""

from __future__ import annotations

import io
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import time
from collections.abc import Iterator
from pathlib import Path

import pikepdf
import pypdf
import pypdfium2
import pytest
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from app.core.config import get_settings
from app.core import document_rendition as rendition
from app.core.document_rendition import (
    ConversionError,
    RenditionOptions,
    SandboxAccount,
    convert_to_pdf,
    find_soffice,
    prepare_document,
    rendition_account,
    run_sandboxed,
    sniff_office_format,
    soffice_command,
    write_profile,
)
from app.core.knowledge_originals import DocumentRef, KnowledgeOriginalStore
from app.main import _restrict_data_directory_permissions
from tests import crafted_documents as crafted

IN_IMAGE = os.environ.get("APERTURE_RENDITION_IMAGE_PROOF") == "1"
ACCOUNT = rendition_account()
SOFFICE = find_soffice()
RUNNABLE = ACCOUNT is not None and SOFFICE is not None and sys.platform == "linux"
SECRET = "SYNTHETICWALLSECRET8842"
ENVIRONMENT_SECRET = os.environ.get("APERTURE_PROOF_SECRET", "")
MAGENTA = (255, 0, 255)
REF = DocumentRef("tenant-wall", "config-wall", "document-wall")
OPTIONS = RenditionOptions(soffice=SOFFICE, timeout_seconds=240.0, memory_mb=4096)

pytestmark = pytest.mark.skipif(
    not RUNNABLE and not IN_IMAGE, reason="needs root, Linux, LibreOffice, and the account"
)


def _as_account(*command: str) -> subprocess.CompletedProcess[str]:
    assert ACCOUNT is not None
    return subprocess.run(
        list(command),
        user=ACCOUNT.uid,
        group=ACCOUNT.gid,
        extra_groups=[],
        env={"PATH": "/usr/bin:/bin"},
        cwd="/",
        capture_output=True,
        text=True,
        timeout=60,
    )


def _secret_pdf() -> bytes:
    document = pikepdf.new()
    font = pikepdf.Dictionary(
        Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1, BaseFont=pikepdf.Name.Helvetica
    )
    page = document.add_blank_page(page_size=(612, 792))
    page.Resources = pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font))
    page.Contents = document.make_stream(
        f"1 0 1 rg 72 200 468 400 re f 0 0 0 rg BT /F1 12 Tf 72 700 Td ({SECRET}) Tj ET".encode()
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _secret_png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (300, 300), MAGENTA).save(buffer, "PNG")
    return buffer.getvalue()


@pytest.fixture
def data_directory() -> Iterator[Path]:
    """The real data directory, restricted the way API startup restricts it."""
    root = Path(get_settings().runtime_state_path).expanduser().parent
    root.mkdir(parents=True, exist_ok=True)
    planted = [root / "wall-proof-secret.txt", root / "wall-proof-secret.pdf", root / "wall-proof-secret.png"]
    planted[0].write_text(f"{SECRET}\n")
    planted[1].write_bytes(_secret_pdf())
    planted[2].write_bytes(_secret_png())
    _restrict_data_directory_permissions()
    yield root
    for path in planted:
        path.unlink(missing_ok=True)


@pytest.fixture
def stand_in_directory() -> Iterator[Path]:
    """Somewhere the account can run a stand-in for soffice from."""
    directory = Path(tempfile.mkdtemp(prefix="wall-stand-in-"))
    directory.chmod(0o755)
    yield directory
    subprocess.run(["rm", "-rf", str(directory)], check=False)


def _stand_in(directory: Path, body: str) -> str:
    """A soffice that drops its evidence into the 'PDF' it writes."""
    script = directory / "soffice"
    script.write_text(
        "#!/bin/sh\n"
        'while [ $# -gt 0 ]; do [ "$1" = "--outdir" ] && { shift; outdir="$1"; }; shift; done\n'
        f"{body}\n"
    )
    script.chmod(0o755)
    return str(script)


def _source(tmp_path: Path, content: bytes) -> tuple[Path, Path]:
    source = tmp_path / "source"
    source.write_bytes(content)
    work = tmp_path / "work"
    work.mkdir()
    return source, work


def test_this_really_is_the_image_and_not_a_skip() -> None:
    assert sys.platform == "linux"
    assert os.geteuid() == 0
    assert ACCOUNT is not None, "the image must have the aperture-render account"
    assert ACCOUNT.uid not in {0, os.getuid()}
    assert ACCOUNT.gid != 0
    assert SOFFICE is not None, "the image must have LibreOffice"
    print(subprocess.run([SOFFICE, "--version"], capture_output=True, text=True).stdout)


def test_the_account_has_no_other_groups_and_no_login() -> None:
    assert ACCOUNT is not None
    assert _as_account("id", "-u").stdout.strip() == str(ACCOUNT.uid)
    assert _as_account("id", "-G").stdout.split() == [str(ACCOUNT.gid)]


def test_the_account_cannot_read_a_secret_in_the_data_directory(data_directory: Path) -> None:
    result = _as_account("cat", str(data_directory / "wall-proof-secret.txt"))
    assert result.returncode != 0
    assert SECRET not in result.stdout


def test_the_account_cannot_list_the_data_directory(data_directory: Path) -> None:
    result = _as_account("ls", str(data_directory))
    assert result.returncode != 0
    assert result.stdout == ""


def test_the_account_cannot_read_stored_originals(data_directory: Path) -> None:
    store = KnowledgeOriginalStore(data_directory / "knowledge_originals")
    store.save(REF, SECRET.encode())
    try:
        original = store.path(REF)
        _restrict_data_directory_permissions()
        result = _as_account("cat", str(original))
        assert result.returncode != 0
        assert SECRET not in result.stdout
    finally:
        store.delete_config(REF.tenant_id, REF.config_id)


def _unreachable(path: Path) -> None:
    """Root reads it, so it is real, and the account is refused, not just missing."""
    assert path.is_file(), f"{path} does not exist, so this would prove nothing"
    assert path.stat().st_size > 0
    result = _as_account("cat", str(path))
    assert result.returncode != 0, f"{path} was readable"
    assert "Permission denied" in result.stderr
    assert result.stdout == ""


@pytest.fixture(scope="module")
def booted() -> dict[str, object]:
    """The real API started on the volume the way the image starts it, on a volume
    an earlier version left world-readable."""
    data = Path(get_settings().runtime_state_path).parent
    data.mkdir(parents=True, exist_ok=True)
    data.chmod(0o755)
    inherited = data / "inherited-from-an-earlier-version.txt"
    inherited.write_text(f"{SECRET}\n")
    inherited.chmod(0o644)
    environment = {
        key: value for key, value in os.environ.items() if key != "PYTEST_CURRENT_TEST"
    }
    environment.update(
        {
            "APERTURE_ENV": "production",
            "APERTURE_SECRET_KEY": "synthetic-wall-proof-signing-key-0123456789abcdef",
            "PYTHONPATH": str(Path(__file__).resolve().parents[1]),
        }
    )
    boot = subprocess.run(
        [sys.executable, str(Path(__file__).with_name("rendition_production_boot.py"))],
        env=environment,
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True,
        text=True,
        timeout=300,
    )
    assert boot.returncode == 0, boot.stderr[-2000:]
    report = json.loads(boot.stdout.strip().splitlines()[-1])
    report["inherited"] = str(inherited)
    report["data"] = str(data)
    return report


def test_the_data_directory_is_a_mounted_volume_as_it_is_in_production(
    booted: dict[str, object],
) -> None:
    assert os.path.ismount(str(booted["data"]))


def test_production_converts_through_the_wall_and_refuses_without_it(
    booted: dict[str, object],
) -> None:
    assert booted["environment"] == "production"
    assert (booted["mode"], booted["notice"], booted["pages"]) == ("pdf", None, 1)
    assert (booted["open_wall_mode"], booted["open_wall_notice"]) == (
        "text",
        "conversion_unavailable",
    )


def test_startup_closes_a_world_readable_volume(booted: dict[str, object]) -> None:
    assert stat.S_IMODE(Path(str(booted["data"])).stat().st_mode) == 0o700
    assert stat.S_IMODE(Path(str(booted["inherited"])).stat().st_mode) == 0o600
    assert booted["umask"] == 0o077


def test_the_account_cannot_read_the_real_databases_or_originals(
    booted: dict[str, object],
) -> None:
    for path in booted["files"]:
        _unreachable(Path(str(path)))


def test_a_world_readable_file_on_the_volume_is_still_out_of_reach(
    booted: dict[str, object],
) -> None:
    """Live volumes hold files written at 0644 before any of this existed. The
    directory above them is what stops the account."""
    data = Path(str(booted["data"]))
    folder = data / "written-by-an-earlier-version"
    folder.mkdir(exist_ok=True)
    folder.chmod(0o755)
    planted = folder / "plain.sqlite3"
    planted.write_text(f"{SECRET}\n")
    planted.chmod(0o644)
    try:
        assert stat.S_IMODE(planted.stat().st_mode) == 0o644
        _unreachable(planted)
    finally:
        planted.unlink()
        folder.rmdir()


@pytest.mark.parametrize("process", ["1", str(os.getpid())], ids=["pid-1", "this-process"])
def test_the_account_cannot_read_the_api_environment_or_memory(process: str) -> None:
    for path in (f"/proc/{process}/environ", f"/proc/{process}/mem"):
        result = _as_account("cat", path)
        assert result.returncode != 0, f"{path} was readable"
        assert not ENVIRONMENT_SECRET or ENVIRONMENT_SECRET not in result.stdout
    assert _as_account("ls", f"/proc/{process}/fd").returncode != 0
    assert _as_account("ls", f"/proc/{process}/root/").returncode != 0


STOCK_PROFILE = "<oor:items xmlns:oor='http://openoffice.org/2001/registry'/>"


def _hostile_run(target: Path, tmp_path: Path, account: SandboxAccount | None) -> Path:
    """Every layer but the OS out of the way: raw file, stock profile."""
    scratch = Path(tempfile.mkdtemp(prefix="wall-run-"))
    try:
        source = scratch / "crafted.docx"
        source.write_bytes(crafted.linked_picture(target.as_uri()))
        write_profile(scratch / "profile", STOCK_PROFILE)
        (scratch / "tmp").mkdir()
        (scratch / "out").mkdir()
        rendition._hand_over(scratch, account)
        assert SOFFICE is not None
        command = soffice_command(SOFFICE, source, sniff_office_format(source), scratch, OPTIONS)
        run_sandboxed(command, scratch, OPTIONS.timeout_seconds, account)
        produced = scratch / "out" / "crafted.pdf"
        assert produced.is_file(), "no PDF means LibreOffice failed, which proves nothing"
        kept = tmp_path / f"rendition-{os.urandom(4).hex()}.pdf"
        kept.write_bytes(produced.read_bytes())
        return kept
    finally:
        shutil.rmtree(scratch, ignore_errors=True)


def _shows_the_secret(pdf: Path) -> bool:
    text = " ".join(page.extract_text() or "" for page in pypdf.PdfReader(pdf).pages)
    document = pypdfium2.PdfDocument(pdf)
    try:
        pages = [document[index].render(scale=0.5).to_pil().convert("RGB") for index in range(len(document))]
    finally:
        document.close()
    painted = any(
        count > 20 and _is_magenta(colour)
        for page in pages
        for count, colour in page.getcolors(1 << 22) or []
    )
    return SECRET in text or painted


def _is_magenta(colour: tuple[int, int, int]) -> bool:
    """Close, not exact: scaling a picture onto a page shifts a channel by a step."""
    return all(abs(channel - wanted) <= 8 for channel, wanted in zip(colour, MAGENTA, strict=True))


def test_the_wall_alone_stops_a_read_every_other_layer_missed(
    data_directory: Path, tmp_path: Path
) -> None:
    target = data_directory / "wall-proof-secret.png"
    assert _shows_the_secret(_hostile_run(target, tmp_path, None)), "the trap must be live as root"
    assert not _shows_the_secret(_hostile_run(target, tmp_path, ACCOUNT))


def test_a_root_run_of_libreoffice_cannot_lock_the_account_out(
    data_directory: Path, tmp_path: Path
) -> None:
    """LibreOffice builds its shared extension cache on first use. Built by a root
    run under the API's umask it is 0700, and every later run as the account aborts,
    so the image has to ship it already built."""
    previous = os.umask(0o077)
    try:
        _hostile_run(data_directory / "wall-proof-secret.png", tmp_path, None)
    finally:
        os.umask(previous)
    cache = Path("/var/spool/libreoffice/uno_packages/cache/uno_packages")
    assert stat.S_IMODE(cache.stat().st_mode) == 0o755
    _hostile_run(data_directory / "wall-proof-secret.png", tmp_path, ACCOUNT)


def _deck() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Ordinary deck words"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _workbook() -> bytes:
    workbook = Workbook()
    workbook.active.append(["Ordinary sheet words"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("name", "content", "words"),
    [
        ("brief.docx", crafted.docx("<w:p><w:r><w:t>Ordinary word words.</w:t></w:r></w:p>"), "Ordinary word words."),
        ("deck.pptx", _deck(), "Ordinary deck words"),
        ("book.xlsx", _workbook(), "Ordinary sheet words"),
    ],
    ids=["docx", "pptx", "xlsx"],
)
def test_ordinary_office_files_convert_with_real_content_under_every_limit(
    tmp_path: Path, name: str, content: bytes, words: str
) -> None:
    originals = KnowledgeOriginalStore(tmp_path / "originals")
    originals.save(REF, content)

    prepared = prepare_document(originals, REF, name=name, options=OPTIONS)

    assert (prepared.mode, prepared.notice) == ("pdf", None)
    assert prepared.page_count == 1
    text = " ".join(page.extract_text() or "" for page in pypdf.PdfReader(prepared.file_path).pages)
    assert words in text


def test_scan_converts_as_the_account(tmp_path: Path) -> None:
    buffer = io.BytesIO()
    Image.new("L", (64, 64), 90).save(buffer, "TIFF")
    originals = KnowledgeOriginalStore(tmp_path / "originals")
    originals.save(REF, buffer.getvalue())
    prepared = prepare_document(originals, REF, name="scan.tif", options=OPTIONS)
    assert (prepared.mode, prepared.page_count) == ("pdf", 1)


def _limit(report: str, name: str) -> str:
    match = re.search(rf"^{name}\s+(\S+)\s+(\S+)", report, re.MULTILINE)
    assert match is not None, name
    assert match.group(1) == match.group(2), f"{name} soft and hard differ"
    return match.group(1)


def test_limits_and_identity_really_apply_inside_the_conversion(
    tmp_path: Path, stand_in_directory: Path
) -> None:
    soffice = _stand_in(
        stand_in_directory,
        '{ printf "%%PDF-1.7\\n"; cat /proc/self/limits; printf "uid=%s\\n" "$(id -u)"; '
        'printf "groups=%s\\n" "$(id -G)"; env; } > "$outdir/input.pdf"',
    )
    source, work = _source(tmp_path, crafted.docx(""))

    produced = convert_to_pdf(source, sniff_office_format(source), work, soffice=soffice, options=OPTIONS)

    report = produced.read_text()
    assert ACCOUNT is not None
    assert f"uid={ACCOUNT.uid}\n" in report
    assert f"groups={ACCOUNT.gid}\n" in report
    assert int(_limit(report, "Max address space")) == 4096 * 1024 * 1024
    assert _limit(report, "Max core file size") == "0"
    assert _limit(report, "Max open files") == "4096"
    assert int(_limit(report, "Max cpu time")) == 240
    assert int(_limit(report, "Max file size")) == 1024 * 1024 * 1024
    assert ENVIRONMENT_SECRET == "" or ENVIRONMENT_SECRET not in report
    assert "APERTURE_" not in report


def test_from_inside_a_conversion_the_secrets_are_refused_and_no_descriptor_leaks(
    data_directory: Path, tmp_path: Path, stand_in_directory: Path
) -> None:
    """Needs no LibreOffice behaviour, so it cannot go quiet on another build."""
    secret = data_directory / "wall-proof-secret.txt"
    soffice = _stand_in(
        stand_in_directory,
        '{ printf "%%PDF-1.7\\n"; '
        f'cat "{secret}" 2>&1; printf "status=%s\\n" "$?"; '
        'cat /proc/1/environ 2>&1; printf "status=%s\\n" "$?"; '
        'ls -l /proc/$$/fd; '
        '} > "$outdir/input.pdf"',
    )
    source, work = _source(tmp_path, crafted.docx(""))

    produced = convert_to_pdf(source, sniff_office_format(source), work, soffice=soffice, options=OPTIONS)

    report = produced.read_text(errors="replace")
    assert SECRET not in report
    assert not ENVIRONMENT_SECRET or ENVIRONMENT_SECRET not in report
    assert "status=0" not in report
    assert report.count("Permission denied") == 2
    descriptors = re.findall(r"^\S+ .* (\d+) -> (.*)$", report, re.MULTILINE)
    assert len(descriptors) >= 3, "the listing must have run"
    # Its stdio, the file it writes, and the shell reading this script.
    allowed = ("/dev/null", "/tmp/aperture-render-", str(stand_in_directory))
    leaked = [target for _number, target in descriptors if not target.startswith(allowed)]
    assert leaked == [], f"the conversion inherited {leaked}"


def test_a_process_that_left_the_session_does_not_outlive_the_conversion(
    tmp_path: Path, stand_in_directory: Path
) -> None:
    soffice = _stand_in(
        stand_in_directory,
        'setsid sh -c \'echo $$ > "$0"; exec sleep 300\' "$outdir/straggler" &\n'
        'until [ -s "$outdir/straggler" ]; do sleep 0.05; done\n'
        '{ printf "%%PDF-1.7\\n"; cat "$outdir/straggler"; } > "$outdir/input.pdf"',
    )
    source, work = _source(tmp_path, crafted.docx(""))

    produced = convert_to_pdf(source, sniff_office_format(source), work, soffice=soffice, options=OPTIONS)

    pid = int(produced.read_text().split()[-1])
    assert _has_exited(pid), "a process that left the session survived the conversion"


def _has_exited(pid: int) -> bool:
    """Gone, or a zombie waiting for PID 1 to reap it: either way it runs no code."""
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        try:
            status = Path(f"/proc/{pid}/status").read_text()
        except FileNotFoundError:
            return True
        if re.search(r"^State:\s+Z", status, re.MULTILINE):
            return True
        time.sleep(0.05)
    return False


@pytest.mark.parametrize(
    "body",
    [
        'ln -s "$SECRET_PDF" "$outdir/input.pdf"',
        'rm -rf "$outdir"; mkdir "$outdir.real"; ln -s "$SECRET_PDF" "$outdir.real/input.pdf"; '
        'ln -s "$outdir.real" "$outdir"',
    ],
    ids=["file-link", "folder-link"],
)
def test_output_cannot_trick_root_into_copying_a_secret(
    data_directory: Path, tmp_path: Path, stand_in_directory: Path, body: str
) -> None:
    soffice = _stand_in(
        stand_in_directory,
        body.replace("$SECRET_PDF", str(data_directory / "wall-proof-secret.pdf")),
    )
    source, work = _source(tmp_path, crafted.docx(""))
    with pytest.raises(ConversionError):
        convert_to_pdf(source, sniff_office_format(source), work, soffice=soffice, options=OPTIONS)
    assert not (work / "rendition.pdf").exists()


def test_scratch_belongs_to_nobody_else_and_leaves_no_trace(
    tmp_path: Path, stand_in_directory: Path
) -> None:
    soffice = _stand_in(
        stand_in_directory,
        '{ printf "%%PDF-1.7\\n"; stat -c "%u %a" "$HOME"; printf "%s\\n" "$HOME"; } > "$outdir/input.pdf"',
    )
    source, work = _source(tmp_path, crafted.docx(""))

    produced = convert_to_pdf(source, sniff_office_format(source), work, soffice=soffice, options=OPTIONS)

    _header, owner_and_mode, home = produced.read_text().splitlines()
    assert ACCOUNT is not None
    assert owner_and_mode == f"{ACCOUNT.uid} 700"
    assert not Path(home).exists()
    assert not stat.S_ISLNK(os.lstat(produced).st_mode)
