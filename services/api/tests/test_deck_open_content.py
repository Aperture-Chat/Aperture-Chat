from __future__ import annotations

import base64
import io

from fastapi.testclient import TestClient
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches

from app.core import deck_template_parse
from app.core.deck_template_parse import parse_deck_template
from app.main import app

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
AUTH = {"x-aperture-user": "user-admin"}


def _png(width: int, height: int, color=(20, 120, 200, 255)) -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGBA", (width, height), color).save(buffer, format="PNG")
    return buffer.getvalue()


def _deck_bytes(*, pictures: int = 1) -> bytes:
    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = "Q3 Review"
    cover.placeholders[1].text = "Board update"

    agenda = presentation.slides.add_slide(presentation.slide_layouts[1])
    agenda.shapes.title.text = "Agenda"
    frame = agenda.placeholders[1].text_frame
    frame.text = "Revenue"
    nested = frame.add_paragraph()
    nested.text = "Up 12%   year over year"
    nested.level = 1
    deep = frame.add_paragraph()
    deep.text = "Deepest"
    deep.level = 4
    agenda.notes_slide.notes_text_frame.text = "Mention the hiring freeze."

    for index in range(pictures):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Chart {index + 1}"
        slide.shapes.add_picture(
            io.BytesIO(_png(800, 450)), Inches(1), Inches(1.5), Inches(6), Inches(3.4)
        )
        # A logo-sized picture is decoration, not the slide's picture.
        slide.shapes.add_picture(io.BytesIO(_png(40, 40)), Inches(0.1), Inches(0.1), Inches(0.3), Inches(0.3))

    compare = presentation.slides.add_slide(presentation.slide_layouts[3])
    compare.shapes.title.text = "Compare"
    compare.placeholders[1].text = "Left side"
    compare.placeholders[2].text = "Right side"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_content_mode_keeps_subtitles_levels_notes_and_columns() -> None:
    result = parse_deck_template("q3.pptx", _deck_bytes(), include_content=True)

    cover, agenda, chart, compare = result.slides
    assert cover.is_title_slide is True
    assert cover.subtitle == "Board update"
    assert cover.bodies == []

    assert agenda.is_title_slide is False
    assert [(p.text, p.level) for p in agenda.bodies[0]] == [
        ("Revenue", 0),
        ("Up 12% year over year", 1),
        ("Deepest", 2),
    ]
    assert agenda.notes == "Mention the hiring freeze."

    assert chart.picture is not None
    assert chart.picture.data_url.startswith("data:image/jpeg;base64,")
    decoded = PILImage.open(io.BytesIO(base64.b64decode(chart.picture.data_url.split(",", 1)[1])))
    assert decoded.size == (800, 450)

    assert [[p.text for p in body] for body in compare.bodies] == [["Left side"], ["Right side"]]
    assert not any("logo" in warning.lower() for warning in result.warnings)


def test_brand_mode_response_is_unchanged() -> None:
    result = parse_deck_template("q3.pptx", _deck_bytes())

    assert all(slide.bodies == [] and slide.picture is None and slide.notes == "" for slide in result.slides)
    assert all(slide.subtitle is None and slide.is_title_slide is False for slide in result.slides)
    assert result.slides[1].blocks == ["Revenue Up 12% year over year Deepest"]


def test_content_mode_pictures_share_one_budget(monkeypatch) -> None:
    monkeypatch.setattr(deck_template_parse, "_MAX_PICTURES_TOTAL_BYTES", 6_000)

    result = parse_deck_template("q3.pptx", _deck_bytes(pictures=3), include_content=True)

    kept = [slide for slide in result.slides if slide.picture is not None]
    assert 0 < len(kept) < 3
    assert sum(1 for warning in result.warnings if "pictures were left out" in warning) == 1


def test_endpoint_returns_content_only_when_asked() -> None:
    client = TestClient(app)
    files = {"file": ("q3.pptx", _deck_bytes(), PPTX_MIME)}

    opened = client.post("/api/drafts/deck-template/parse?content=true", files=files, headers=AUTH)
    assert opened.status_code == 200
    assert opened.json()["slides"][1]["notes"] == "Mention the hiring freeze."

    brand = client.post(
        "/api/drafts/deck-template/parse",
        files={"file": ("q3.pptx", _deck_bytes(), PPTX_MIME)},
        headers=AUTH,
    )
    assert brand.status_code == 200
    assert brand.json()["slides"][1]["bodies"] == []
